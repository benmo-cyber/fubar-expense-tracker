from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
import csv
import io
import re
import uuid
from fastapi import APIRouter, Depends, File, HTTPException, UploadFile
from fastapi.responses import Response
from openpyxl import load_workbook
from pydantic import BaseModel
from sqlalchemy import and_, or_
from sqlalchemy.orm import Session
from app.core.database import get_db
from app.core.security import get_password_hash
from app.services.passwords import generate_temp_password, must_change_after_issue
from app.middleware.auth import get_current_admin_user, get_current_user
from app.models import (
    Expense, ExpenseReport, GLAccount, Merchant, Notice, ReportStatus, Trip, User, UserRole,
)
from app.services.corrections import can_send_back, can_withdraw, report_is_open, stamp_lines
from app.services.gl_rollup import UNASSIGNED, rollup
from app.services.receipt_pack import bundle_name, grouped_receipts, pack_report, receipt_folder_name
from app.services.reports import can_view_report, refresh_report, team_member_ids
from app.services.storage import storage_service

router = APIRouter(tags=["Workspace"])
TEMPLATE_PATH = Path(__file__).resolve().parents[2] / "templates" / "expense_report_template.xlsx"
EN_DASH = "\u2013"


class PersonIn(BaseModel):
    email: str
    full_name: str
    role: str = "sales"
    supervisor_id: uuid.UUID | None = None


class PersonPatch(BaseModel):
    full_name: str | None = None
    role: str | None = None
    is_active: bool | None = None
    supervisor_id: uuid.UUID | None = None


class ReportIn(BaseModel):
    period_start: date
    period_end: date
    title: str | None = None


class TripIn(BaseModel):
    name: str


class RejectIn(BaseModel):
    notes: str


def is_admin(user: User) -> bool:
    return user.role == UserRole.ADMIN


def chart_rows(db: Session) -> list[dict]:
    return [
        {"id": str(user.id), "supervisor_id": str(user.supervisor_id) if user.supervisor_id else None}
        for user in db.query(User.id, User.supervisor_id).all()
    ]


def viewer_can_open(user: User, report: ExpenseReport, db: Session) -> bool:
    status = report.status.value if hasattr(report.status, "value") else str(report.status)
    return can_view_report(str(user.id), is_admin(user), str(report.user_id), status, chart_rows(db))


def person_dict(user: User) -> dict:
    return {
        "id": str(user.id),
        "email": user.email,
        "full_name": user.full_name,
        "role": user.role.value if hasattr(user.role, "value") else str(user.role),
        "is_active": bool(user.is_active),
        "is_superuser": bool(user.is_superuser),
        "supervisor_id": str(user.supervisor_id) if user.supervisor_id else None,
    }


def report_dict(report: ExpenseReport, db: Session) -> dict:
    expenses = db.query(Expense).filter(
        Expense.report_id == report.id,
        Expense.removed_at.is_(None),
    ).all()
    by_gl: dict[str, Decimal] = {}
    by_merchant: dict[str, Decimal] = {}
    for expense in expenses:
        gl_name = expense.gl_account.account_name if expense.gl_account else "Unassigned"
        merchant_name = expense.merchant.name if expense.merchant else (expense.merchant_name or "Unknown")
        by_gl[gl_name] = by_gl.get(gl_name, Decimal("0")) + expense.amount
        by_merchant[merchant_name] = by_merchant.get(merchant_name, Decimal("0")) + expense.amount
    trips = []
    for trip in report.trips:
        trip_expenses = [expense for expense in expenses if expense.trip_id == trip.id]
        trip_gl: dict[str, Decimal] = {}
        trip_merchant: dict[str, Decimal] = {}
        for expense in trip_expenses:
            gl_name = expense.gl_account.account_name if expense.gl_account else "Unassigned"
            merchant_name = expense.merchant.name if expense.merchant else (expense.merchant_name or "Unknown")
            trip_gl[gl_name] = trip_gl.get(gl_name, Decimal("0")) + expense.amount
            trip_merchant[merchant_name] = trip_merchant.get(merchant_name, Decimal("0")) + expense.amount
        trips.append({
            "id": str(trip.id),
            "name": trip.name,
            "total": float(sum((expense.amount for expense in trip_expenses), Decimal("0"))),
            "by_gl": [{"name": name, "amount": float(amount)} for name, amount in trip_gl.items()],
            "by_merchant": [{"name": name, "amount": float(amount)} for name, amount in trip_merchant.items()],
        })
    status = report.status.value if hasattr(report.status, "value") else str(report.status)
    total = sum((expense.amount or 0) for expense in expenses)
    gl_groups = rollup(
        [
            (str(expense.gl_account_id) if expense.gl_account_id else UNASSIGNED, expense.amount or Decimal("0"))
            for expense in expenses
        ],
        db.query(GLAccount).all(),
    )
    report.total_amount = total or Decimal("0")
    report.expense_count = len(expenses)
    return {
        "id": str(report.id),
        "user_id": str(report.user_id),
        "user_name": report.user.full_name if report.user else "",
        "title": report.title or f"{report.period_start} to {report.period_end}",
        "period_start": report.period_start.isoformat(),
        "period_end": report.period_end.isoformat(),
        "status": status.lower() if isinstance(status, str) else status,
        "total": float(total or 0),
        "expense_count": len(expenses),
        "review_notes": report.review_notes,
        "screenshot_url": report.screenshot_url,
        "by_gl": [{"name": name, "amount": float(amount)} for name, amount in by_gl.items()],
        "gl_groups": gl_groups,
        "by_merchant": [{"name": name, "amount": float(amount)} for name, amount in by_merchant.items()],
        "trips": trips,
        "expenses": [
            {
                "id": str(expense.id),
                "merchant_name": expense.merchant.name if expense.merchant else (expense.merchant_name or "Unknown"),
                "amount": float(expense.amount),
                "expense_date": expense.expense_date.isoformat(),
                "trip_id": str(expense.trip_id) if expense.trip_id else None,
                "trip_name": next((trip.name for trip in report.trips if trip.id == expense.trip_id), None),
                "category_name": expense.category.name if expense.category else None,
                "category_id": str(expense.category_id) if expense.category_id else None,
                "description": expense.description,
                "gl_name": expense.gl_account.account_name if expense.gl_account else None,
                "gl_code": expense.gl_account.account_code if expense.gl_account else None,
                "parent_code": expense.gl_account.parent.account_code if expense.gl_account and expense.gl_account.parent else None,
                "parent_name": expense.gl_account.parent.account_name if expense.gl_account and expense.gl_account.parent else None,
                "receipt_url": expense.receipt_url,
                "notes": expense.notes,
            }
            for expense in expenses
        ],
    }


def sync_reminders(db: Session) -> None:
    today = date.today()
    open_reports = db.query(ExpenseReport).filter(
        ExpenseReport.period_end < today,
        ExpenseReport.status.in_([ReportStatus.DRAFT, ReportStatus.REJECTED]),
    ).all()
    for report in open_reports:
        owner = report.user
        if owner is None:
            continue
        label = report.title or "expense report"
        messages = [(owner.id, f"Your {label} is still open. Submit it for review.")]
        if owner.supervisor_id:
            messages.append((owner.supervisor_id, f"{owner.full_name} has an open expense report."))
        for user_id, message in messages:
            existing = db.query(Notice).filter(
                Notice.user_id == user_id,
                Notice.report_id == report.id,
                Notice.read_at.is_(None),
            ).first()
            if existing is None:
                db.add(Notice(user_id=user_id, report_id=report.id, message=message))
    db.commit()


@router.get("/summary")
def summary(user_id: uuid.UUID | None = None, current_user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    target_id = user_id if user_id and is_admin(current_user) else current_user.id
    expenses = db.query(Expense).filter(
        Expense.user_id == target_id,
        Expense.removed_at.is_(None),
    ).all()
    by_gl: dict[str, Decimal] = {}
    by_merchant: dict[str, Decimal] = {}
    for expense in expenses:
        gl_name = expense.gl_account.account_name if expense.gl_account else "Unassigned"
        merchant_name = expense.merchant.name if expense.merchant else (expense.merchant_name or "Unknown")
        by_gl[gl_name] = by_gl.get(gl_name, Decimal("0")) + (expense.amount or 0)
        by_merchant[merchant_name] = by_merchant.get(merchant_name, Decimal("0")) + (expense.amount or 0)
    reports = db.query(ExpenseReport).filter(ExpenseReport.user_id == target_id).all()
    return {
        "user_id": str(target_id),
        "spend": float(sum((expense.amount or 0) for expense in expenses)),
        "by_gl": [{"name": name, "amount": float(amount)} for name, amount in sorted(by_gl.items())],
        "by_merchant": [{"name": name, "amount": float(amount)} for name, amount in sorted(by_merchant.items(), key=lambda item: item[1], reverse=True)],
        "open_reports": sum(1 for report in reports if report.status in (ReportStatus.DRAFT, ReportStatus.REJECTED)),
        "awaiting_review": sum(1 for report in reports if report.status == ReportStatus.SUBMITTED),
    }


@router.get("/people")
def list_people(current_user: User = Depends(get_current_admin_user), db: Session = Depends(get_db)):
    return [person_dict(user) for user in db.query(User).order_by(User.full_name).all()]


@router.post("/people")
def invite_person(body: PersonIn, current_user: User = Depends(get_current_admin_user), db: Session = Depends(get_db)):
    if "@" not in body.email:
        raise HTTPException(status_code=400, detail="Enter a valid email")
    if db.query(User).filter(User.email == body.email.lower()).first():
        raise HTTPException(status_code=400, detail="That email is already in use")
    role = UserRole.ADMIN if body.role == "admin" else UserRole.SALES
    supervisor_id = body.supervisor_id or current_user.id
    temporary_password = generate_temp_password()
    user = User(
        email=body.email.lower(),
        full_name=body.full_name.strip(),
        hashed_password=get_password_hash(temporary_password),
        role=role,
        is_active=True,
        must_change_password=True,
        supervisor_id=supervisor_id,
    )
    db.add(user)
    db.commit()
    db.refresh(user)
    payload = person_dict(user)
    payload["temporary_password"] = temporary_password
    return payload


@router.patch("/people/{user_id}")
def update_person(user_id: uuid.UUID, body: PersonPatch, current_user: User = Depends(get_current_admin_user), db: Session = Depends(get_db)):
    user = db.query(User).filter(User.id == user_id).first()
    if user is None:
        raise HTTPException(status_code=404, detail="Person not found")
    if body.is_active is False and user.id == current_user.id:
        raise HTTPException(status_code=400, detail="You cannot deactivate yourself")
    if body.full_name is not None:
        user.full_name = body.full_name.strip()
    if body.role is not None:
        user.role = UserRole.ADMIN if body.role == "admin" else UserRole.SALES
    if body.is_active is not None:
        user.is_active = body.is_active
    if body.supervisor_id is not None:
        user.supervisor_id = body.supervisor_id
    db.commit()
    db.refresh(user)
    return person_dict(user)


@router.post("/people/{user_id}/temporary-password")
def issue_temporary_password(user_id: uuid.UUID, current_user: User = Depends(get_current_admin_user), db: Session = Depends(get_db)):
    user = db.query(User).filter(User.id == user_id).first()
    if user is None:
        raise HTTPException(status_code=404, detail="Person not found")
    if not user.is_active:
        raise HTTPException(status_code=400, detail="Reactivate this person before issuing a password")
    temporary_password = generate_temp_password()
    user.hashed_password = get_password_hash(temporary_password)
    user.must_change_password = must_change_after_issue(current_user.id, user.id)
    db.commit()
    return {"temporary_password": temporary_password, "must_change_password": bool(user.must_change_password)}


@router.get("/merchants")
def list_merchants(current_user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    query = db.query(Expense).filter(Expense.removed_at.is_(None))
    if not is_admin(current_user):
        query = query.filter(Expense.user_id == current_user.id)
    totals: dict[str, dict] = {}
    for expense in query.all():
        name = expense.merchant.name if expense.merchant else (expense.merchant_name or "Unknown")
        key = str(expense.merchant_id or name)
        bucket = totals.setdefault(key, {"id": str(expense.merchant_id) if expense.merchant_id else None, "name": name, "amount": Decimal("0"), "count": 0})
        bucket["amount"] += expense.amount or 0
        bucket["count"] += 1
    rows = sorted(totals.values(), key=lambda item: item["amount"], reverse=True)
    return [{"id": row["id"], "name": row["name"], "amount": float(row["amount"]), "count": row["count"]} for row in rows]


@router.post("/merchants/{merchant_id}/merge")
def merge_merchants(merchant_id: uuid.UUID, into_id: uuid.UUID, current_user: User = Depends(get_current_admin_user), db: Session = Depends(get_db)):
    source = db.query(Merchant).filter(Merchant.id == merchant_id).first()
    target = db.query(Merchant).filter(Merchant.id == into_id).first()
    if source is None or target is None or source.id == target.id:
        raise HTTPException(status_code=400, detail="Choose two different merchants")
    for expense in db.query(Expense).filter(Expense.merchant_id == source.id).all():
        expense.merchant_id = target.id
    for alias in list(source.aliases):
        alias.merchant_id = target.id
    db.delete(source)
    db.commit()
    return {"merged_into": str(target.id)}


@router.get("/reports")
def list_reports(user_id: uuid.UUID | None = None, current_user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    query = db.query(ExpenseReport)
    if is_admin(current_user) and user_id:
        query = query.filter(ExpenseReport.user_id == user_id)
    elif not is_admin(current_user):
        team = team_member_ids(str(current_user.id), chart_rows(db))
        visible = [ExpenseReport.user_id == current_user.id]
        if team:
            visible.append(and_(
                ExpenseReport.user_id.in_(list(team)),
                ExpenseReport.status.in_([ReportStatus.DRAFT, ReportStatus.SUBMITTED, ReportStatus.REJECTED]),
            ))
        query = query.filter(or_(*visible))
    reports = query.order_by(ExpenseReport.period_start.desc()).all()
    payload = [report_dict(report, db) for report in reports]
    db.commit()
    return payload


@router.post("/reports")
def create_report(body: ReportIn, current_user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    if body.period_end < body.period_start:
        raise HTTPException(status_code=400, detail="The end date is before the start date")
    report = ExpenseReport(
        user_id=current_user.id,
        period_start=body.period_start,
        period_end=body.period_end,
        title=body.title or f"{body.period_start.isoformat()} to {body.period_end.isoformat()}",
        total_amount=0,
        currency="USD",
        expense_count=0,
        status=ReportStatus.DRAFT,
    )
    db.add(report)
    db.commit()
    db.refresh(report)
    return report_dict(report, db)


@router.get("/reports/{report_id}")
def get_report(report_id: uuid.UUID, current_user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    report = db.query(ExpenseReport).filter(ExpenseReport.id == report_id).first()
    if report is None:
        raise HTTPException(status_code=404, detail="Report not found")
    if not viewer_can_open(current_user, report, db):
        raise HTTPException(status_code=403, detail="Not authorized")
    payload = report_dict(report, db)
    db.commit()
    return payload


def _sorted_expenses(payload: dict) -> list[dict]:
    return sorted(payload["expenses"], key=lambda item: (item["expense_date"], item.get("merchant_name") or ""))


def _gl_label(expense: dict) -> str:
    code = (expense.get("gl_code") or "").strip()
    name = (expense.get("gl_name") or "").strip()
    if code and name:
        return f"{code} {EN_DASH} {name}"
    return name or code


def _line_description(expense: dict) -> str:
    text = (expense.get("description") or "").strip()
    if text:
        return text
    return (expense.get("category_name") or "").strip()


def _fill_workbook(payload: dict):
    workbook = load_workbook(TEMPLATE_PATH)
    sheet = workbook["Expense Report"]
    for row in range(2, 30):
        for col in range(1, 10):
            sheet.cell(row, col).value = None
    expenses = _sorted_expenses(payload)
    for index, expense in enumerate(expenses):
        row = 2 + index
        sheet.cell(row, 1).value = date.fromisoformat(expense["expense_date"])
        sheet.cell(row, 1).number_format = "MM/DD/YYYY"
        sheet.cell(row, 2).value = expense["merchant_name"]
        sheet.cell(row, 3).value = _line_description(expense)
        sheet.cell(row, 5).value = _gl_label(expense)
        sheet.cell(row, 6).value = expense["amount"]
        sheet.cell(row, 6).number_format = '#,##0.00'
        sheet.cell(row, 7).value = "Yes" if expense.get("receipt_url") else "No"
        sheet.cell(row, 9).value = expense.get("notes") or ""
    last = max(29, 1 + len(expenses))
    sheet["N1"] = payload["user_name"]
    sheet["N2"] = date.fromisoformat(payload["period_end"])
    sheet["N2"].number_format = "MM/DD/YYYY"
    sheet["N3"] = f"=SUM(F2:F{last})+Mileage!D33"
    return workbook


def _account_listing_rows(expenses: list[dict]) -> list[list]:
    workbook = load_workbook(TEMPLATE_PATH, data_only=False)
    listing = workbook["Account Listing"]
    totals: dict[str, float] = {}
    for expense in expenses:
        label = _gl_label(expense)
        totals[label] = totals.get(label, 0) + float(expense["amount"])
    rows = []
    matched = set()
    for row in range(2, 37):
        number = listing.cell(row, 1).value
        name = listing.cell(row, 2).value
        label = listing.cell(row, 3).value
        if number is None and name is None:
            continue
        if isinstance(number, float) and number.is_integer():
            number = int(number)
        amount = totals.get(str(label), 0) if label else 0
        if label in totals:
            matched.add(str(label))
        rows.append([number, name, label, round(amount, 2)])
    for label, amount in totals.items():
        if label not in matched:
            rows.append(["", "", label, round(amount, 2)])
    return rows


def _csv_body(payload: dict) -> str:
    expenses = _sorted_expenses(payload)
    total = round(sum(float(expense["amount"]) for expense in expenses), 2)
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow(["Name", payload["user_name"]])
    writer.writerow(["Date", payload["period_end"]])
    writer.writerow(["Total Owed", total])
    writer.writerow([])
    writer.writerow(["Expense Report"])
    writer.writerow([
        "Date",
        "Vendor / Merchant",
        "Description",
        "Attendees (if Applicable)",
        "GL Account",
        "Amount",
        "Do you have a reciept? ",
        "Is this tied to a specific PO/SO? What number?",
        "Notes",
    ])
    for expense in expenses:
        writer.writerow([
            expense["expense_date"],
            expense["merchant_name"],
            _line_description(expense),
            "",
            _gl_label(expense),
            expense["amount"],
            "Yes" if expense.get("receipt_url") else "No",
            "",
            expense.get("notes") or "",
        ])
    writer.writerow([])
    writer.writerow(["Mileage"])
    writer.writerow(["Day", "Mileage Total", "Comments"])
    for day in range(1, 32):
        writer.writerow([day, "", ""])
    writer.writerow(["Total miles", 0, "Rate", 0.7, "Mileage amount", 0])
    writer.writerow([])
    writer.writerow(["Account Listing"])
    writer.writerow(["Account Number", "Account Name", "GL (Number - Name)", "Amount"])
    writer.writerows(_account_listing_rows(expenses))
    return buffer.getvalue()


def _receipt_file(url: str) -> bytes | None:
    getter = getattr(storage_service, "get_file_path", None)
    if getter is None:
        return None
    path = getter(url)
    if not path.is_file():
        return None
    return path.read_bytes()


def _download_name(payload: dict, extension: str) -> str:
    raw = f"FUBAR-{(payload['user_name'] or 'report')}-{payload['title']}"
    safe = re.sub(r"[^A-Za-z0-9._-]+", "-", raw).strip("-")
    return f"{safe}.{extension}"


@router.get("/reports/{report_id}/export")
def export_report(
    report_id: uuid.UUID,
    file_format: str = "xlsx",
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    report = db.query(ExpenseReport).filter(ExpenseReport.id == report_id).first()
    if report is None:
        raise HTTPException(status_code=404, detail="Report not found")
    if not viewer_can_open(current_user, report, db):
        raise HTTPException(status_code=403, detail="Not authorized")
    payload = report_dict(report, db)
    db.commit()
    if file_format == "csv":
        return Response(
            _csv_body(payload),
            media_type="text/csv",
            headers={"Content-Disposition": f'attachment; filename="{_download_name(payload, "csv")}"'},
        )
    workbook = _fill_workbook(payload)
    output = io.BytesIO()
    workbook.save(output)
    return Response(
        output.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{_download_name(payload, "xlsx")}"'},
    )


@router.get("/reports/{report_id}/receipts")
def export_receipt_pack(
    report_id: uuid.UUID,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    report = db.query(ExpenseReport).filter(ExpenseReport.id == report_id).first()
    if report is None:
        raise HTTPException(status_code=404, detail="Report not found")
    if not viewer_can_open(current_user, report, db):
        raise HTTPException(status_code=403, detail="Not authorized")
    payload = report_dict(report, db)
    db.commit()
    workbook = _fill_workbook(payload)
    output = io.BytesIO()
    workbook.save(output)
    folders = []
    for folder, named in grouped_receipts(payload["user_name"], payload["title"], _sorted_expenses(payload)):
        files = []
        for filename, expense in named:
            content = _receipt_file(expense["receipt_url"])
            if content is not None:
                files.append((filename, content))
        if files:
            folders.append((folder, files))
    sheet_name = f"{bundle_name(payload['user_name'], payload['title'])}.xlsx"
    zip_name = f"{receipt_folder_name(payload['user_name'], payload['title'])}.zip"
    return Response(
        pack_report(sheet_name, output.getvalue(), folders),
        media_type="application/zip",
        headers={"Content-Disposition": f'attachment; filename="{zip_name}"'},
    )


@router.post("/reports/{report_id}/trips")
def add_trip(report_id: uuid.UUID, body: TripIn, current_user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    report = db.query(ExpenseReport).filter(ExpenseReport.id == report_id).first()
    if report is None:
        raise HTTPException(status_code=404, detail="Report not found")
    if report.user_id != current_user.id and not is_admin(current_user):
        raise HTTPException(status_code=403, detail="Not authorized")
    if not report_is_open(report.status):
        raise HTTPException(status_code=400, detail="Pull this report back before adding a trip.")
    if not body.name.strip():
        raise HTTPException(status_code=400, detail="Enter a trip name.")
    trip = Trip(report_id=report.id, name=body.name.strip())
    db.add(trip)
    db.commit()
    db.refresh(report)
    return report_dict(report, db)


def _owned_report(db: Session, report_id: uuid.UUID, user: User, admin_ok: bool = False) -> ExpenseReport:
    report = db.query(ExpenseReport).filter(ExpenseReport.id == report_id).first()
    if report is None:
        raise HTTPException(status_code=404, detail="Report not found")
    if admin_ok and is_admin(user):
        return report
    if report.user_id != user.id:
        raise HTTPException(status_code=403, detail="Not authorized")
    return report


@router.post("/reports/{report_id}/submit")
def submit_report(report_id: uuid.UUID, current_user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    report = _owned_report(db, report_id, current_user)
    if report.status not in (ReportStatus.DRAFT, ReportStatus.REJECTED):
        raise HTTPException(status_code=400, detail="This report is not open for submission")
    refresh_report(db, report)
    if not report.expense_count:
        raise HTTPException(status_code=400, detail="Add at least one expense before submitting")
    report.status = ReportStatus.SUBMITTED
    report.submitted_at = datetime.utcnow()
    report.review_notes = None
    stamp_lines(db, report)
    db.commit()
    db.refresh(report)
    return report_dict(report, db)


@router.post("/reports/{report_id}/withdraw")
def withdraw_report(report_id: uuid.UUID, current_user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    report = _owned_report(db, report_id, current_user)
    if not can_withdraw(report.status):
        raise HTTPException(status_code=400, detail="Only a submitted report can be pulled back")
    report.status = ReportStatus.DRAFT
    report.submitted_at = None
    stamp_lines(db, report)
    db.commit()
    db.refresh(report)
    return report_dict(report, db)


@router.post("/reports/{report_id}/approve")
def approve_report(report_id: uuid.UUID, current_user: User = Depends(get_current_admin_user), db: Session = Depends(get_db)):
    report = db.query(ExpenseReport).filter(ExpenseReport.id == report_id).first()
    if report is None:
        raise HTTPException(status_code=404, detail="Report not found")
    if report.status != ReportStatus.SUBMITTED:
        raise HTTPException(status_code=400, detail="Only a submitted report can be approved")
    report.status = ReportStatus.APPROVED
    report.reviewed_by = current_user.id
    report.finalized_at = datetime.utcnow()
    stamp_lines(db, report, current_user.id)
    db.commit()
    db.refresh(report)
    return report_dict(report, db)


@router.post("/reports/{report_id}/reject")
def reject_report(report_id: uuid.UUID, body: RejectIn, current_user: User = Depends(get_current_admin_user), db: Session = Depends(get_db)):
    notes = body.notes.strip()
    if not notes:
        raise HTTPException(status_code=400, detail="Add a note about what to change")
    report = db.query(ExpenseReport).filter(ExpenseReport.id == report_id).first()
    if report is None:
        raise HTTPException(status_code=404, detail="Report not found")
    if not can_send_back(report.status):
        raise HTTPException(status_code=400, detail="Only a submitted or approved report can be sent back")
    report.status = ReportStatus.REJECTED
    report.review_notes = notes
    report.reviewed_by = current_user.id
    stamp_lines(db, report)
    if report.user_id:
        db.add(Notice(
            user_id=report.user_id,
            report_id=report.id,
            message=f"Your report was sent back: {notes}",
        ))
    db.commit()
    db.refresh(report)
    return report_dict(report, db)


@router.post("/reports/{report_id}/screenshot")
def upload_screenshot(
    report_id: uuid.UUID,
    file: UploadFile = File(...),
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    report = _owned_report(db, report_id, current_user, admin_ok=True)
    if not (file.content_type or "").startswith("image/"):
        raise HTTPException(status_code=400, detail="Upload an image")
    report.screenshot_url = storage_service.upload_file(file.file, file.filename or "screenshot.jpg", file.content_type or "image/jpeg")
    db.commit()
    db.refresh(report)
    return report_dict(report, db)


@router.get("/notices")
def list_notices(current_user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    sync_reminders(db)
    rows = db.query(Notice).filter(Notice.user_id == current_user.id).order_by(Notice.created_at.desc()).limit(20).all()
    return [{
        "id": str(row.id),
        "message": row.message,
        "report_id": str(row.report_id) if row.report_id else None,
        "read": row.read_at is not None,
    } for row in rows]


@router.post("/notices/{notice_id}/read")
def read_notice(notice_id: uuid.UUID, current_user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    notice = db.query(Notice).filter(Notice.id == notice_id, Notice.user_id == current_user.id).first()
    if notice is None:
        raise HTTPException(status_code=404, detail="Notice not found")
    notice.read_at = datetime.utcnow()
    db.commit()
    return {"ok": True}
